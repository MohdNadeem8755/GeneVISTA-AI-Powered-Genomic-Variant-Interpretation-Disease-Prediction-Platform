"""Import published TP53 tables; verify source allele/protein mappings locally.

Read-only XLSX extraction needs openpyxl (build-time only). Runtime uses gzip/JSON.
Run: python -m src.build_tp53_evidence --source-dir data/incoming/tp53_evidence
"""
import argparse
from collections import Counter, defaultdict
from datetime import datetime, timezone
import gzip
import hashlib
import json
from pathlib import Path
import re
import sqlite3

from .tp53_evidence import AA, ARTICLE, CSPEC, RULE_TRANSCRIPT, RULE_VERSION, TRANSCRIPT, assess_variant

ROOT = Path(__file__).resolve().parents[1]
BASE = 'https://media.springernature.com/original/springer-static/esm/art%3A10.1186%2Fs13073-025-01536-3/MediaObjects/'
NCBI = 'https://eutils.ncbi.nlm.nih.gov/entrez/eutils/efetch.fcgi?db=nuccore&id=NM_000546.5,NM_000546.6&rettype=gb&retmode=text'
BASES = 'TCAG'
CODONS = dict(zip((a+b+c for a in BASES for b in BASES for c in BASES),
                  'FFLLSSSSYY**CC*WLLLLPPPPHHQQRRRRIIIMTTTTNNKKSSRRVVVVAAAADDEEGGGG'))


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def verify_transcripts(text):
    sequences = {}
    for record in text.split('//'):
        version = re.search(r'^VERSION\s+(NM_000546\.\d+)', record, re.M)
        if not version:
            continue
        cds = re.findall(r'^     CDS\s+(\d+)\.\.(\d+)\s*$', record, re.M)
        if len(cds) != 1 or '\nORIGIN' not in record:
            raise ValueError('Expected one exact, contiguous CDS in each NCBI transcript.')
        sequence = re.sub('[^acgt]', '', record.split('\nORIGIN', 1)[1]).upper()
        start, end = map(int, cds[0])
        coding = sequence[start-1:end]
        if len(coding) != 1182 or end > len(sequence):
            raise ValueError('Unexpected TP53 CDS length.')
        protein = ''.join(CODONS[coding[i:i+3]] for i in range(0, len(coding), 3))
        if protein[-1] != '*' or '*' in protein[:-1]:
            raise ValueError('Unexpected TP53 translation.')
        if version[1] in sequences:
            raise ValueError('Duplicate transcript accession.')
        sequences[version[1]] = {'cds': coding, 'protein': protein[:-1], 'cds_start': start, 'cds_end': end}
    if set(sequences) != {TRANSCRIPT, RULE_TRANSCRIPT}:
        raise ValueError('Both exact transcript versions are required.')
    if sequences[TRANSCRIPT]['cds'] != sequences[RULE_TRANSCRIPT]['cds']:
        raise ValueError('Transcript coding sequences differ. Manual mapping is required.')
    proof = {'accessions': [RULE_TRANSCRIPT, TRANSCRIPT], 'cds_identical': True,
             'cds_length': 1182, 'cds_sha256': hashlib.sha256(sequences[TRANSCRIPT]['cds'].encode()).hexdigest(),
             'scope': 'Coding SNVs only. No equivalence is asserted for UTRs, genomic coordinates or other transcripts.',
             'source_url': NCBI}
    return sequences[TRANSCRIPT], proof


def table_rows(path, sheet_name, expected_headers):
    from openpyxl import load_workbook
    # A successful HTTP response can still be an HTML challenge, not a workbook.
    if path.read_bytes()[:4] != b'PK\x03\x04':
        raise ValueError(f'{path.name} is not an XLSX archive.')
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = wb[sheet_name]
        headers = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))
        if tuple(str(v).strip() for v in headers[:len(expected_headers)]) != tuple(expected_headers):
            raise ValueError(f'Unexpected headers in {path.name}.')
        for number, row in enumerate(sheet.iter_rows(min_row=4, values_only=True), 4):
            if any(v is not None for v in row):
                yield number, row
    finally:
        wb.close()


def normalize_code(value):
    codes = {'No evidence', 'BP4', 'BP4_Moderate', 'PP3', 'PP3_Moderate', 'BS3', 'BS3_Supporting', 'PS3', 'PS3_Moderate', 'PS3_Supporting'}
    result = str(value).replace('_moderate', '_Moderate')
    if result not in codes:
        raise ValueError(f'Unexpected source criterion: {result}')
    return result


def build(source_dir):
    transcript, proof = verify_transcripts((source_dir / 'NM_000546_versions.gb').read_text(encoding='utf-8'))
    tables = {n: source_dir / f'13073_2025_1536_MOESM{n}_ESM.xlsx' for n in [2, 3]}
    computational, functional = {}, defaultdict(list)
    excluded_rows = []
    for number, row in table_rows(tables[2], 'Table S2', ['Transcript change (NM_000546.6)', 'Protein change (NP_000537.3)', 'Align-GVGD', 'BayesDel', 'Preliminary bioinformatic code (missense only)', 'Maximum SpliceAI delta score (raw, 10000 bp window)']):
        change, protein, align, bayes, code, splice = row[:6]
        cm = re.fullmatch(r'c\.([1-9]\d*)([ACGT])>([ACGT])', change or '')
        pm = re.fullmatch(r'p\.([A-Z][a-z]{2})([1-9]\d*)([A-Z][a-z]{2})', protein or '')
        if not cm or not pm or pm[1] not in AA or pm[3] not in AA:
            raise ValueError(f'Unexpected variant at S2 row {number}.')
        position, ref, alt = int(cm[1]), cm[2], cm[3]
        aa_position = (position-1)//3 + 1
        if position > 1179 or transcript['cds'][position-1] != ref or int(pm[2]) != aa_position:
            raise ValueError(f'CDS reference mismatch at S2 row {number}.')
        codon = list(transcript['cds'][(aa_position-1)*3:aa_position*3])
        original = CODONS[''.join(codon)]
        codon[(position-1)%3] = alt
        if original != AA[pm[1]] or CODONS[''.join(codon)] != AA[pm[3]] or ref == alt:
            raise ValueError(f'Protein mapping mismatch at S2 row {number}.')
        if change in computational:
            raise ValueError(f'Duplicate coding allele {change}.')
        if not isinstance(splice, (int, float)) or not 0 <= splice <= 1 or not isinstance(bayes, (int,float)):
            raise ValueError(f'Missing/invalid predictors at S2 row {number}.')
        computational[change] = {'protein_change': f'{AA[pm[1]]}{pm[2]}{AA[pm[3]]}',
            'align_gvgd': align.removeprefix('Class '), 'bayesdel': bayes, 'spliceai': splice,
            'published_code': normalize_code(code), 'source_row': number, 'source_sheet': 'Table S2',
            'splice_note': row[6]}
    expected = ['Protein change', 'Seq_change (specified when Kotler or Funk result changes for the same missense variant)',
                'Kato class (PMID: 12826609)', 'Funk class (PMID 39774325)', 'Giacomelli class (PMID: 30224644)',
                'Kotler class (PMID: 29979965)', 'Kawaguchi class (PMID: 16007150)', 'Preliminary functional code (missense only)']
    for number, row in table_rows(tables[3], 'Table S3', expected):
        protein, selector = row[:2]
        pm = re.fullmatch(r'([ARNDCQEGHILKMFPSTWYV])([1-9]\d*)([ARNDCQEGHILKMFPSTWYV])', protein or '')
        if not pm or not 1 <= int(pm[2]) <= 393:
            raise ValueError(f'Protein reference mismatch at S3 row {number}.')
        if transcript['protein'][int(pm[2])-1] != pm[1]:
            excluded_rows.append({'sheet': 'Table S3', 'row': number, 'protein_change': protein,
                                  'reason': 'Reference amino acid differs from NP_000537.3; assay background needs review.'})
            continue
        functional[protein].append({'sequence_selector': selector, 'published_code': normalize_code(row[7]),
            'assays': dict(zip(['Kato', 'Funk', 'Giacomelli', 'Kotler', 'Kawaguchi'], row[2:7])),
            'source_row': number, 'source_sheet': 'Table S3'})
    sources = [{'label': 'TP53 study and supplementary tables (2025)', 'url': ARTICLE},
               {'label': 'Current TP53 specification v2.4', 'url': CSPEC}]
    files = {p.name: {'sha256': sha256(p), 'bytes': p.stat().st_size, 'url': BASE+p.name} for p in tables.values()}
    for name, url in [('NM_000546_versions.gb', NCBI), ('cspec_GN009.html', CSPEC)]:
        path = source_dir/name
        files[name] = {'sha256': sha256(path), 'bytes': path.stat().st_size, 'url': url}
    # Refuse an unverified version instead of labeling arbitrary downloaded HTML.
    if 'TP53 Version 2.4' not in (source_dir/'cspec_GN009.html').read_text(encoding='utf-8'):
        raise ValueError('Cached CSpec v2.4 was not verified.')
    metadata = {'schema_version': 1, 'table_transcript': TRANSCRIPT, 'protein_accession': 'NP_000537.3',
        'rule_version': RULE_VERSION, 'functional_codes_version': 'Published v2 tables; not recalculated for v2.4',
        'transcript_verification': proof, 'sources': sources, 'files': files,
        'attribution': 'Fortuno et al., Genome Medicine (2025), DOI 10.1186/s13073-025-01536-3, Tables S2 and S3. Converted to JSON; allele matching and current computational candidates added by GeneVISTA.',
        'license': 'CC BY 4.0', 'license_url': 'https://creativecommons.org/licenses/by/4.0/',
        'created_utc': datetime.now(timezone.utc).isoformat(),
        'excluded': 'Table S1 was not imported: download returned an HTML challenge. Canonical splice, indel and truncating variants are outside this pilot.',
        'excluded_functional_rows': excluded_rows,
        'counts': {'coding_alleles': len(computational), 'functional_rows': sum(map(len, functional.values())), 'protein_changes': len(functional), 'excluded_functional_rows': len(excluded_rows)}}
    return {'metadata': metadata, 'computational': computational, 'functional': dict(functional)}


def coverage(bundle, database):
    counts, examples = Counter(), {}
    with sqlite3.connect(database.resolve().as_uri()+'?mode=ro', uri=True) as con:
        con.row_factory = sqlite3.Row
        for row in con.execute("SELECT variation_id,name,gene_symbol FROM variants WHERE gene_symbol='TP53'"):
            result = assess_variant(dict(row), bundle)
            counts[result['status']] += 1
            if result['status'] == 'available':
                fc = result['functional']
                counts['functional_'+fc['status']] += 1
                counts['criterion_'+fc['criterion_status']] += 1
                if result['conflicting_evidence_directions']:
                    counts['conflicting_evidence_directions'] += 1
                for key in ('available', fc['criterion_status']):
                    examples.setdefault(key, {'variation_id': row['variation_id'], 'name': row['name']})
    return {'counts': dict(counts), 'examples': examples,
            'interpretation': 'Coverage only. No clinical labels or test outcomes were used; these counts are not accuracy.'}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-dir', type=Path, default=ROOT/'data/incoming/tp53_evidence')
    parser.add_argument('--output', type=Path, default=ROOT/'data/processed/tp53_evidence_v1.json.gz')
    parser.add_argument('--database', type=Path, default=ROOT/'data/app/genevista.sqlite')
    args = parser.parse_args()
    bundle = build(args.source_dir)
    report = {'metadata': bundle['metadata'], 'coverage': coverage(bundle, args.database)}
    payload = json.dumps(bundle, ensure_ascii=False, allow_nan=False, sort_keys=True, separators=(',', ':')).encode()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    # Atomic replacement: failed imports never replace a working package.
    temporary = args.output.with_suffix('.tmp')
    temporary.write_bytes(gzip.compress(payload, mtime=0))
    temporary.replace(args.output)
    report['artifact_sha256'] = sha256(args.output)
    report['artifact_bytes'] = args.output.stat().st_size
    report_path = ROOT/'reports/tp53_evidence_v1/import_report.json'
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(report, indent=2, allow_nan=False)+'\n', encoding='utf-8')
    print(json.dumps({'counts': bundle['metadata']['counts'], 'coverage': report['coverage'],
                      'artifact_bytes': report['artifact_bytes']}, indent=2))


if __name__ == '__main__':
    main()
